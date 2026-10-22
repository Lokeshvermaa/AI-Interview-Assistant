"""Reads the uploaded resume/JD files (PDF, DOCX, TXT) and returns plain text."""

import io

from docx import Document
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


def extract_text(filename: str, data: bytes) -> str:
    # Pick the reader based on the file extension.
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if ext == ".pdf":
        return _from_pdf(data)
    if ext == ".docx":
        return _from_docx(data)
    if ext in {".txt", ".md"}:
        return data.decode("utf-8", errors="ignore").strip()

    raise ValueError(
        f"Unsupported file type '{ext or filename}'. "
        f"Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
    )


def _from_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(pages).strip()


def _from_docx(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts).strip()
